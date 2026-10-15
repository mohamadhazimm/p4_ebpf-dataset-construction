#!/usr/bin/env python3
"""Generate a presentation from the committed 224-sample evaluation artifacts."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
COHORT = ROOT / "data/malware/revision4/family-heldout-balanced-224-v1-20261014"
REPORTS = COHORT / "reports/balanced_9network_9host"
SCORES_PATH = REPORTS / "category_analysis/overall_and_category_scores.csv"
METRICS_PATH = REPORTS / "cv_metrics.csv"
FEATURES_PATH = ROOT / "data/malware/revision4/family-heldout-stable-151-v1-20261014/feature_breadth_comparison/balanced_9network_9host_feature_columns.json"
INVENTORY_PATH = ROOT / "data/malware/malware_sample_inventory_all_expansions_20261014.csv"

WIDTH = 13.333
HEIGHT = 7.5
FONT = "DejaVu Sans"
BG = "F4F4EF"
PAPER = "FFFFFF"
INK = "1F302E"
MUTED = "61716D"
GRID = "D9E0DB"
GREEN = "2D745C"
TEAL = "418C98"
CORAL = "D76751"
GOLD = "C99537"
PALE_GREEN = "E2EEE8"
PALE_TEAL = "E2EFF0"
PALE_CORAL = "F5E5E0"
PALE_GOLD = "F3EDD9"
MODEL_COLORS = {"network": TEAL, "host": GREEN, "hybrid": CORAL}
MODEL_LABELS = {"network": "Network", "host": "Host / eBPF", "hybrid": "Hybrid"}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as source:
        return list(csv.DictReader(source))


def add_shape(slide, x, y, width, height, fill, rounded=False, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.08
    return shape


def add_text(
    slide,
    x,
    y,
    width,
    height,
    text,
    size=16,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.03,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(str(text).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(3)
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = rgb(color)
    return shape


def add_rule(slide, x1, y1, x2, y2, color=GRID, width=1):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def add_notes(slide, notes: str) -> None:
    slide.notes_slide.notes_text_frame.text = notes


def add_header(slide, section: str, title: str, slide_no: int, source: str):
    add_text(slide, 0.66, 0.34, 9.8, 0.28, section.upper(), 9, GREEN, True)
    add_text(slide, 0.66, 0.72, 11.95, 0.55, title, 25, INK, True)
    add_rule(slide, 0.66, 1.43, 12.66, 1.43, GRID, 1)
    add_rule(slide, 0.66, 7.06, 12.66, 7.06, GRID, 0.7)
    add_text(slide, 0.68, 7.10, 11.2, 0.19, source, 7.3, MUTED)
    add_text(slide, 12.02, 7.08, 0.58, 0.2, f"{slide_no:02d}", 8, MUTED, align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, width, height, title, body, accent=GREEN, fill=PAPER):
    add_shape(slide, x, y, width, height, fill, rounded=True)
    add_shape(slide, x, y, 0.07, height, accent)
    add_text(slide, x + 0.23, y + 0.18, width - 0.42, 0.34, title, 13, accent, True)
    add_text(slide, x + 0.23, y + 0.60, width - 0.42, height - 0.76, body, 12, INK)


def add_metric_tile(slide, x, y, width, value, label, accent):
    add_shape(slide, x, y, width, 1.02, PAPER, rounded=True)
    add_shape(slide, x, y, 0.07, 1.02, accent)
    add_text(slide, x + 0.2, y + 0.1, width - 0.32, 0.46, value, 23, accent, True)
    add_text(slide, x + 0.2, y + 0.59, width - 0.32, 0.27, label, 10, MUTED, True)


def add_base_slide(prs, section, title, slide_no, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    add_header(slide, section, title, slide_no, source)
    return slide


def percentage(value: float) -> str:
    return f"{value:.1%}"


def make_cover(prs, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(INK)
    add_shape(slide, 9.72, 0, 3.61, HEIGHT, GREEN)
    add_shape(slide, 10.24, 0.78, 2.38, 1.25, TEAL, rounded=True)
    add_shape(slide, 10.82, 2.35, 1.80, 1.25, GOLD, rounded=True)
    add_shape(slide, 9.95, 3.92, 2.67, 1.25, CORAL, rounded=True)
    add_text(slide, 0.82, 0.74, 8.2, 0.35, "RESEARCH PRESENTATION  /  15 OCTOBER 2026", 10, "BFD1C7", True)
    add_text(slide, 0.82, 1.58, 8.25, 1.52, "Hybrid Malware\nDetection", 34, PAPER, True)
    add_text(slide, 0.85, 3.45, 7.75, 0.88, "Comparing network telemetry, host eBPF, and their combination", 19, "E0E8E3")
    add_rule(slide, 0.85, 5.65, 8.1, 5.65, "5D766D", 1)
    add_text(slide, 0.85, 5.95, 7.8, 0.62, "224 malware samples  ·  138 benign controls  ·  grouped out-of-fold evaluation", 12, "BFD1C7")
    add_text(slide, 0.85, 6.85, 7.9, 0.28, "P4 + eBPF research project  |  Dataset study", 9, "AABCB4")
    add_notes(slide, notes)
    return slide


def add_category_chart(slide, category_scores):
    labels = {
        "backdoor": "Backdoor",
        "botnet": "Botnet",
        "ddos": "DDoS",
        "other_elf": "Other ELF",
        "ransomware": "Ransomware",
        "trojan": "Trojan*",
    }
    x_label, x_count, x_bar, bar_width = 0.83, 2.28, 3.08, 7.4
    axis_min, axis_max = 0.70, 1.00
    row_y, row_gap = 2.03, 0.64
    for tick in (0.70, 0.80, 0.90, 1.00):
        position = x_bar + (tick - axis_min) / (axis_max - axis_min) * bar_width
        add_rule(slide, position, 1.82, position, 5.80, GRID, 0.6)
        add_text(slide, position - 0.25, 1.60, 0.5, 0.17, f"{tick:.0%}", 8, MUTED, align=PP_ALIGN.CENTER)
    for index, category in enumerate(("backdoor", "botnet", "ddos", "other_elf", "ransomware", "trojan")):
        y = row_y + index * row_gap
        records = category_scores[category]
        sample_count = records["network"]["malware_samples"]
        add_text(slide, x_label, y + 0.04, 1.37, 0.30, labels[category], 11, INK, True)
        count_color = CORAL if category == "trojan" else MUTED
        count_text = f"n={sample_count}" + (" (small)" if category == "trojan" else "")
        add_text(slide, x_count, y + 0.04, 0.78, 0.30, count_text, 8.2, count_color, True)
        for offset, feature_set in enumerate(("network", "host", "hybrid")):
            record = records[feature_set]
            value = float(record["balanced_accuracy"])
            width = max(0.08, (value - axis_min) / (axis_max - axis_min) * bar_width)
            bar_y = y + offset * 0.15
            add_shape(slide, x_bar, bar_y, bar_width, 0.09, "E9ECE8")
            add_shape(slide, x_bar, bar_y, width, 0.09, MODEL_COLORS[feature_set])
            value_x = min(x_bar + bar_width + 0.07, x_bar + width + 0.07)
            add_text(slide, value_x, bar_y - 0.045, 0.47, 0.18, f"{value:.2f}", 7.5, MODEL_COLORS[feature_set], True)
    legend_y = 6.17
    for index, feature_set in enumerate(("network", "host", "hybrid")):
        x = 4.1 + index * 1.72
        add_shape(slide, x, legend_y, 0.16, 0.12, MODEL_COLORS[feature_set])
        add_text(slide, x + 0.22, legend_y - 0.03, 1.38, 0.2, MODEL_LABELS[feature_set], 8.5, MUTED)


def create_deck(output: Path) -> Path:
    metric_rows = read_csv(METRICS_PATH)
    metrics = {row["feature_set"]: row for row in metric_rows}
    score_rows = read_csv(SCORES_PATH)
    category_scores: dict[str, dict[str, dict[str, str]]] = {}
    overall_scores = {}
    for row in score_rows:
        if row["score_scope"] == "overall":
            overall_scores[row["feature_set"]] = row
        elif row["score_scope"] == "attack_category_vs_benign":
            category_scores.setdefault(row["category"], {})[row["feature_set"]] = row
    feature_sets = json.loads(FEATURES_PATH.read_text(encoding="utf-8"))
    network_feature_count = len(feature_sets["network"])
    host_feature_count = len(feature_sets["host"])
    hybrid_feature_count = len(feature_sets["hybrid"])
    inventory = read_csv(INVENTORY_PATH)
    acquired = len(inventory)
    modeled_malware = sum(
        row.get("included_in_expanded_family_grouped_model", "").casefold() == "yes"
        for row in inventory
    )
    no_telemetry = acquired - modeled_malware

    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)
    prs.core_properties.title = "Hybrid Malware Detection: Network, Host/eBPF, and Combined Models"
    prs.core_properties.subject = "Professor presentation based on the 224-malware grouped evaluation"
    prs.core_properties.author = "Hybrid Malware Detection Study"
    prs.core_properties.keywords = "malware detection, eBPF, network telemetry, grouped cross-validation"
    notes_for_markdown = []

    notes = (
        "Open with the question: does combining host and network observations help distinguish this malware cohort from a controlled benign workload set? "
        "The headline result is a modest hybrid improvement over host-only under grouped out-of-fold evaluation. This is not a live P4 hardware result or a claim of universal superiority."
    )
    make_cover(prs, notes)
    notes_for_markdown.append(("1. Hybrid Malware Detection", notes))

    notes = (
        "The experiment compares three binary classifiers on identical samples and grouped folds. Network-only uses packet and flow counters, host-only uses eBPF-derived system-call counters, and hybrid combines them. "
        "The question is whether the two sensor views add complementary evidence under comparable feature budgets, not whether the model identifies a malware family."
    )
    slide = add_base_slide(prs, "01 / Question", "Do two sensor views add useful evidence?", 2, "Source: docs/STUDY_WALKTHROUGH.md")
    add_card(slide, 0.82, 1.86, 3.72, 2.20, "NETWORK", "What crossed the interface?\n\nPacket counts, bytes, rates, and flow summaries.", TEAL, PALE_TEAL)
    add_card(slide, 4.80, 1.86, 3.72, 2.20, "HOST / eBPF", "What did the process do?\n\nSelected socket, open, and write syscall counters.", GREEN, PALE_GREEN)
    add_card(slide, 8.78, 1.86, 3.72, 2.20, "HYBRID", "Does the joint view help?\n\nBoth sensor sets, evaluated on the same samples.", CORAL, PALE_CORAL)
    add_text(slide, 0.92, 4.68, 11.5, 0.60, "Primary outcome: malware-versus-benign discrimination, not family classification.", 18, INK, True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        1.1,
        5.48,
        11.1,
        0.48,
        f"Same cohort  ·  Same grouped folds  ·  {network_feature_count} / {host_feature_count} / {hybrid_feature_count} model features",
        12,
        MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_notes(slide, notes)
    notes_for_markdown.append(("2. Research question", notes))

    notes = (
        f"The inventory contains {acquired} acquisitions. {modeled_malware} malware samples had archived telemetry and were included; {no_telemetry} did not and were excluded. "
        "The final binary dataset adds 138 benign workloads, yielding 362 evaluated samples. The newest batch had 42 acquisitions, 38 telemetry captures, and four without usable telemetry. "
        "The 195 telemetry windows and 21,189 eBPF events refer to the newest batch's successful captures, not the whole cohort."
    )
    slide = add_base_slide(prs, "02 / Cohort", "From acquisitions to the evaluated cohort", 3, "Source: malware_sample_inventory_all_expansions_20261014.csv; current feature dataset")
    add_metric_tile(slide, 0.88, 1.84, 2.62, str(acquired), "MALWARE ACQUISITIONS", GOLD)
    add_metric_tile(slide, 3.65, 1.84, 2.62, str(modeled_malware), "MALWARE MODELED", GREEN)
    add_metric_tile(slide, 6.42, 1.84, 2.62, str(no_telemetry), "NO USABLE TELEMETRY", CORAL)
    add_metric_tile(slide, 9.19, 1.84, 3.10, "138", "BENIGN CONTROLS", TEAL)
    add_shape(slide, 1.02, 3.35, 5.16, 1.46, PALE_GREEN, rounded=True)
    add_text(slide, 1.33, 3.63, 4.56, 0.39, "224 malware samples", 23, GREEN, True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.33, 4.12, 4.56, 0.31, "archived telemetry; included in training/evaluation", 10.5, MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, 6.28, 3.71, 0.69, 0.46, "+", 28, MUTED, True, align=PP_ALIGN.CENTER)
    add_shape(slide, 7.03, 3.35, 5.16, 1.46, PALE_TEAL, rounded=True)
    add_text(slide, 7.34, 3.63, 4.56, 0.39, "138 benign workloads", 23, TEAL, True, align=PP_ALIGN.CENTER)
    add_text(slide, 7.34, 4.12, 4.56, 0.31, "28 package-family groups", 10.5, MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, 1.0, 5.36, 11.2, 0.45, "362 total evaluated samples  |  Newest 42-sample batch: 38 captures, 4 collection failures", 12.5, INK, True, align=PP_ALIGN.CENTER)
    add_notes(slide, notes)
    notes_for_markdown.append(("3. Cohort and exclusions", notes))

    notes = (
        "Candidate archives were selected from MalwareBazaar metadata, verified by SHA-256, staged in a guest restored from a clean snapshot, and run one at a time. "
        "The malware VM and sink VM used only the internal malnet network. Packet capture and eBPF telemetry were collected in parallel. The 20-second cap and 0.5-second windows were consistent across the collection. "
        "A nonzero exit or timeout did not change the malware label; missing archived telemetry did exclude a sample."
    )
    slide = add_base_slide(prs, "03 / Collection", "Controlled execution, two synchronized sensors", 4, "Source: scripts/stage_malware_revision4.sh; scripts/vbox_malware_batch.sh; docs/STUDY_WALKTHROUGH.md")
    steps = [
        ("01", "Acquire", "Metadata-selected\nLinux ELF candidates", GOLD),
        ("02", "Verify", "Encrypted ZIP\nSHA-256 check", TEAL),
        ("03", "Isolate", "Clean VM snapshot\ninternal malnet only", GREEN),
        ("04", "Observe", "PCAP + eBPF\n0.5 s windows", CORAL),
        ("05", "Aggregate", "One feature row\nper sample", GOLD),
    ]
    for index, (number, title, body, accent) in enumerate(steps):
        x = 0.82 + index * 2.48
        add_shape(slide, x, 2.12, 2.15, 2.12, PAPER, rounded=True)
        add_text(slide, x + 0.16, 2.31, 0.54, 0.28, number, 10, accent, True)
        add_text(slide, x + 0.16, 2.77, 1.85, 0.35, title, 16, INK, True)
        add_text(slide, x + 0.16, 3.30, 1.86, 0.64, body, 11, MUTED)
        if index < len(steps) - 1:
            add_text(slide, x + 2.17, 2.92, 0.3, 0.3, ">", 16, MUTED, True, align=PP_ALIGN.CENTER)
    add_shape(slide, 1.0, 4.76, 11.3, 0.86, PALE_CORAL, rounded=True)
    add_text(slide, 1.25, 4.98, 10.8, 0.40, "Safety boundary: no NAT or public network interface on the malware VM", 14, CORAL, True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.2, 5.92, 10.9, 0.34, "20-second execution limit  ·  Separate clean restore per sample  ·  Sink VM on internal network", 10.5, MUTED, align=PP_ALIGN.CENTER)
    add_notes(slide, notes)
    notes_for_markdown.append(("4. Isolated collection", notes))

    notes = (
        "The raw collection produces multiple time windows per sample. Dataset construction aggregates each selected base signal using mean, maximum, sum, standard deviation, and active-window count. "
        "Nine network signal families produce 45 values; nine selected host/eBPF families also produce 45. Hybrid concatenates them for 90. "
        "The feature contract intentionally excludes CPU and memory observer metrics, timestamps, labels, filenames, process details, and execution status."
    )
    slide = add_base_slide(prs, "04 / Features", "Equal sensor budgets; shared sample-level summaries", 5, "Source: src/build_dataset.py; feature_breadth_comparison/balanced_9network_9host_feature_columns.json")
    add_text(slide, 0.93, 1.80, 2.82, 0.32, "BASE SIGNAL FAMILIES", 9, MUTED, True)
    add_card(slide, 0.87, 2.24, 3.65, 2.24, f"NETWORK  ·  {network_feature_count // 5} SIGNALS", "Packets, bytes, flow count\nSent/received counts\nPacket and byte rates", TEAL, PALE_TEAL)
    add_card(slide, 4.84, 2.24, 3.65, 2.24, f"HOST / eBPF  ·  {host_feature_count // 5} SIGNALS", "Socket/connect/sendto\nSuccessful open calls\nSuccessful write calls/bytes", GREEN, PALE_GREEN)
    add_card(slide, 8.81, 2.24, 3.65, 2.24, "SUMMARY EACH SIGNAL", "Mean  ·  Max  ·  Sum\nStandard deviation\nActive-window count", GOLD, PALE_GOLD)
    add_text(slide, 0.95, 5.03, 3.4, 0.55, f"{network_feature_count // 5} × 5 = {network_feature_count}", 23, TEAL, True, align=PP_ALIGN.CENTER)
    add_text(slide, 4.96, 5.03, 3.4, 0.55, f"{host_feature_count // 5} × 5 = {host_feature_count}", 23, GREEN, True, align=PP_ALIGN.CENTER)
    add_text(slide, 9.04, 5.03, 3.1, 0.55, f"{network_feature_count} + {host_feature_count} = {hybrid_feature_count}", 23, CORAL, True, align=PP_ALIGN.CENTER)
    add_text(slide, 0.95, 5.83, 11.3, 0.43, "Syscall write counters are not proof of regular-file writes; no file-descriptor/path attribution is available.", 10.4, CORAL, True, align=PP_ALIGN.CENTER)
    add_notes(slide, notes)
    notes_for_markdown.append(("5. Feature construction", notes))

    notes = (
        "StratifiedGroupKFold uses five folds and seed 42. The 57 groups are 29 malware signature groups plus 28 benign package-family groups; no group crosses folds. "
        "Within each fold, median imputation is fit only on the training partition, followed by the same 500-tree Random Forest configuration. "
        "Each sample receives one held-out probability for each feature set. Balanced accuracy is the mean of sensitivity and specificity, which is useful when class sizes differ."
    )
    slide = add_base_slide(prs, "05 / Evaluation", "Keep related samples together; test on held-out groups", 6, "Source: src/train_compare.py; cv_metrics.json")
    add_metric_tile(slide, 0.92, 1.86, 2.55, "5-fold", "STRATIFIED GROUP CV", GREEN)
    add_metric_tile(slide, 3.62, 1.86, 2.55, "57", "INDEPENDENT GROUPS", TEAL)
    add_metric_tile(slide, 6.32, 1.86, 2.55, "500", "TREES / MODEL / FOLD", GOLD)
    add_metric_tile(slide, 9.02, 1.86, 3.30, "0", "GROUPS CROSSING FOLDS", CORAL)
    add_card(slide, 0.94, 3.40, 5.42, 2.07, "GROUPING", "Malware: same source signature stays in one fold.\nBenign: same package family stays in one fold.", GREEN, PALE_GREEN)
    add_card(slide, 6.78, 3.40, 5.42, 2.07, "MODEL PIPELINE", "Fold-local median imputation → Random Forest.\nSame splits and settings for all three feature sets.", TEAL, PALE_TEAL)
    add_text(slide, 1.0, 5.90, 11.4, 0.37, "Report pooled out-of-fold metrics from held-out predictions; never score a model on its training rows.", 11.2, INK, True, align=PP_ALIGN.CENTER)
    add_notes(slide, notes)
    notes_for_markdown.append(("6. Evaluation protocol", notes))

    notes = (
        "These pooled out-of-fold values come directly from cv_metrics.csv. Hybrid balanced accuracy is 0.929, compared with 0.918 for host-only and 0.874 for network-only. "
        "Hybrid ROC-AUC is 0.985. Balanced accuracy is not the same as accuracy: it gives equal weight to sensitivity and specificity. "
        "The displayed bar axis starts at 0.70 to make small differences readable, so compare the printed values rather than visual bar length alone."
    )
    slide = add_base_slide(prs, "06 / Results", "Hybrid leads; the gain over host-only is modest", 7, "Source: current cohort cv_metrics.csv; 362 pooled OOF predictions per model")
    add_text(slide, 0.93, 1.70, 4.2, 0.28, "POOLED OUT-OF-FOLD SCORE", 9, MUTED, True)
    chart_x, chart_width = 4.35, 6.25
    min_value, max_value = 0.70, 1.00
    for tick in (0.70, 0.80, 0.90, 1.00):
        x = chart_x + (tick - min_value) / (max_value - min_value) * chart_width
        add_rule(slide, x, 2.05, x, 5.75, GRID, 0.6)
        add_text(slide, x - 0.25, 1.84, 0.5, 0.16, f"{tick:.0%}", 8, MUTED, align=PP_ALIGN.CENTER)
    for index, name in enumerate(("network", "host", "hybrid")):
        y = 2.27 + index * 1.06
        metric = metrics[name]
        add_text(slide, 0.92, y + 0.25, 1.75, 0.32, MODEL_LABELS[name], 13, MODEL_COLORS[name], True)
        for row_offset, key, label in ((0.05, "balanced_accuracy", "Balanced accuracy"), (0.48, "roc_auc", "ROC-AUC")):
            value = float(metric[key])
            bar_y = y + row_offset
            add_text(slide, 2.95, bar_y - 0.02, 1.38, 0.18, label, 8.3, MUTED)
            width = max(0.05, (value - min_value) / (max_value - min_value) * chart_width)
            add_shape(slide, chart_x, bar_y, chart_width, 0.17, "E7EBE7")
            add_shape(slide, chart_x, bar_y, width, 0.17, MODEL_COLORS[name])
            add_text(slide, chart_x + chart_width + 0.08, bar_y - 0.03, 0.56, 0.20, percentage(value), 9.5, MODEL_COLORS[name], True)
    add_shape(slide, 1.02, 5.74, 11.1, 0.60, INK, rounded=True)
    add_text(slide, 1.30, 5.86, 10.55, 0.30, "Hybrid vs host: +1.1 pp balanced accuracy  ·  5 fewer malware misses  ·  same 9 false positives", 12, PAPER, True, align=PP_ALIGN.CENTER)
    add_notes(slide, notes)
    notes_for_markdown.append(("7. Overall results", notes))

    notes = (
        "This chart uses balanced accuracy for each metadata category versus the same 138 benign controls, from the global model's out-of-fold predictions. "
        "The sample count is shown beside each category because the groups vary from 2 Trojan samples to 86 botnet samples. "
        "Hybrid is strongest for botnet and DDoS, and ties host for backdoor and the two-sample Trojan slice. Network scores higher on the small Other ELF slice, while host is strongest for ransomware. "
        "These are descriptive slices and not independently trained category classifiers."
    )
    slide = add_base_slide(prs, "07 / Category slices", "Category results vary; sample counts matter", 8, "Source: category_analysis/overall_and_category_scores.csv; metadata labels")
    add_text(slide, 0.86, 1.62, 11.6, 0.30, "Balanced accuracy: one metadata category vs the same benign controls  ·  values are not multiclass scores", 9.4, MUTED)
    add_category_chart(slide, category_scores)
    add_shape(slide, 0.88, 6.48, 11.6, 0.36, PALE_GOLD, rounded=True)
    add_text(slide, 1.04, 6.54, 11.25, 0.21, "* Trojan n=2: too small for a stable performance estimate. Categories are MalwareBazaar metadata.", 8.6, INK, True, align=PP_ALIGN.CENTER)
    add_notes(slide, notes)
    notes_for_markdown.append(("8. Category results", notes))

    notes = (
        "The confusion counts make the operating-point tradeoff concrete. Network has 32 false negatives and 15 false positives. Host has 22 and 9. Hybrid has 17 and 9. "
        "Relative to host-only, the hybrid model reduces false negatives by five without increasing false positives on this pooled OOF set. "
        "This is a cohort-specific threshold result at 0.5, not a guarantee at another deployment threshold."
    )
    slide = add_base_slide(prs, "08 / Error profile", "Hybrid reduced misses without adding false alarms here", 9, "Source: current cohort cv_metrics.csv; threshold 0.5")
    columns = [("network", 0.93), ("host", 4.81), ("hybrid", 8.69)]
    max_count = 35
    for name, x in columns:
        metric = metrics[name]
        add_shape(slide, x, 2.02, 3.48, 2.88, PAPER, rounded=True)
        add_text(slide, x + 0.18, 2.24, 3.1, 0.34, MODEL_LABELS[name], 16, MODEL_COLORS[name], True, align=PP_ALIGN.CENTER)
        for index, (field, label, accent) in enumerate((("fn", "Missed malware (FN)", CORAL), ("fp", "Benign flagged (FP)", GOLD))):
            y = 3.0 + index * 0.87
            value = int(metric[field])
            add_text(slide, x + 0.20, y, 2.96, 0.21, f"{label}   {value}", 10, INK, True)
            add_shape(slide, x + 0.20, y + 0.30, 2.98, 0.16, "E7EBE7")
            add_shape(slide, x + 0.20, y + 0.30, 2.98 * value / max_count, 0.16, accent)
    add_shape(slide, 1.00, 5.36, 11.27, 0.78, PALE_GREEN, rounded=True)
    add_text(slide, 1.25, 5.55, 10.75, 0.36, "Hybrid vs host-only: 5 more malware detections, with false positives unchanged at 9.", 15, GREEN, True, align=PP_ALIGN.CENTER)
    add_notes(slide, notes)
    notes_for_markdown.append(("9. Error profile", notes))

    notes = (
        "The evidence is limited by the size and provenance of the cohort, the generated benign workload mix, and metadata-derived malware categories. "
        "The study does not yet validate network features against a live P4 target and should not be presented as field performance. "
        "The next experiment should increase independent malware signatures and benign environments, verify behaviors where possible, and compare packet-capture-derived network counters with a running P4 pipeline."
    )
    slide = add_base_slide(prs, "09 / Discussion", "What this result supports—and what it does not", 10, "Source: docs/STUDY_WALKTHROUGH.md, limitations")
    add_card(slide, 0.91, 1.92, 5.47, 2.04, "SUPPORTED BY THIS RUN", "A grouped, equal-budget comparison on 224 malware and 138 generated benign samples.\nA modest hybrid improvement over host-only.", GREEN, PALE_GREEN)
    add_card(slide, 6.79, 1.92, 5.47, 2.04, "NOT ESTABLISHED", "General superiority of hybrid detection.\nReal-world false-positive rates.\nLive P4 hardware performance.", CORAL, PALE_CORAL)
    add_text(slide, 0.95, 4.48, 11.45, 0.30, "Important limitations", 13, INK, True)
    bullets = [
        "Metadata labels are not independently verified behavior; signatures can mix categories.",
        "Generated controls do not represent all ordinary Linux activity; Trojan n=2 is especially tentative.",
        "eBPF writes are syscall-level; network features are packet-capture aggregates, not P4-target measurements.",
    ]
    for index, bullet in enumerate(bullets):
        add_shape(slide, 1.03, 4.97 + index * 0.48, 0.12, 0.12, (TEAL, GOLD, CORAL)[index])
        add_text(slide, 1.28, 4.91 + index * 0.48, 10.75, 0.31, bullet, 10.8, MUTED)
    add_notes(slide, notes)
    notes_for_markdown.append(("10. Limitations and next steps", notes))

    notes = (
        "The committed feature matrix is sufficient to rerun training, out-of-fold scoring, and category analysis without rerunning malware. "
        "The walkthrough document is the detailed narrative and includes exact commands, caveats, and a full artifact map. "
        "Raw malware, packet captures, per-run VM outputs, and detailed event-level collector streams are excluded. "
        "Close with the proposed next steps and invite questions about group splits, label provenance, and the distinction between PCAP aggregates and live P4 measurements."
    )
    slide = add_base_slide(prs, "10 / Reproduction", "The model evaluation is reproducible from committed inputs", 11, "Source: docs/STUDY_WALKTHROUGH.md; repository artifact map")
    add_card(slide, 0.94, 1.90, 5.45, 2.22, "START HERE", "docs/STUDY_WALKTHROUGH.md\n\nChronological method, full results, limitations, commands, and this presentation's story.", GREEN, PALE_GREEN)
    add_card(slide, 6.80, 1.90, 5.45, 2.22, "REPRODUCE", "Committed engineered 224-cohort CSV\n→ train_compare.py\n→ analyze_malware_category_results.py", TEAL, PALE_TEAL)
    add_text(slide, 1.02, 4.56, 11.15, 0.32, "For questions", 13, INK, True)
    add_text(slide, 1.03, 4.99, 11.12, 0.92, "Why group by signature/package?  ·  What does balanced accuracy add?  ·  Why is the hybrid gain modest?\nWhat would change with real benign traffic or live P4 counters?  ·  How should metadata label uncertainty be handled?", 12, MUTED)
    add_shape(slide, 1.01, 6.18, 11.23, 0.42, INK, rounded=True)
    add_text(slide, 1.22, 6.25, 10.80, 0.22, "Takeaway: promising complementarity in this cohort; broader, independent validation is the next scientific step.", 10.4, PAPER, True, align=PP_ALIGN.CENTER)
    add_notes(slide, notes)
    notes_for_markdown.append(("11. Reproduction and discussion", notes))

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    notes_path = output.with_name("speaker_notes.md")
    notes_content = ["# Speaker Notes: Hybrid Malware Detection", "", "Companion to `malware_detection_study.pptx`.", ""]
    for heading, note in notes_for_markdown:
        notes_content.extend((f"## {heading}", "", note, ""))
    notes_path.write_text("\n".join(notes_content), encoding="utf-8")
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "presentations/malware_detection_study.pptx",
    )
    args = parser.parse_args()
    result = create_deck(args.output.expanduser().resolve())
    print(f"PowerPoint: {result}")
    print(f"Speaker notes: {result.with_name('speaker_notes.md')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())